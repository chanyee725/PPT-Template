"""Rough slide previews for when LibreOffice is not installed.

    python tools/preview.py out/interim-report.pptx out/preview

Draws shape boxes, fills and wrapped text, with charts as empty frames. It
ignores alignment and vertical centring, so use it to catch overlaps and
crowding, not to judge the final look. Writes slide-NN.png and sheet-N.png
(six slides per sheet).
"""

import re
import subprocess
import sys
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Emu

DPI = 72
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"


def font_path(pattern: str) -> str:
    return subprocess.run(
        ["fc-match", "-f", "%{file}", pattern], capture_output=True, text=True
    ).stdout


SANS = font_path("Pretendard")
BOLD = font_path("Pretendard:bold")


def scheme(deck) -> dict:
    colors = {}
    for rel in deck.slide_masters[0].part.rels.values():
        if rel.reltype.endswith("/theme"):
            xml = rel.target_part.blob.decode("utf-8")
            colors = dict(re.findall(r'<a:(\w+)><a:srgbClr val="(\w{6})"/>', xml))
    alias = {"tx1": "dk1", "bg1": "lt1", "tx2": "dk2", "bg2": "lt2"}
    out = {k: colors.get(v, "FF00FF") for k, v in alias.items()}
    out.update(colors)
    return out


def colour(el, sch):
    if el is None:
        return None
    for c in el.iter():
        if c.tag == A + "srgbClr":
            return "#" + c.get("val")
        if c.tag == A + "schemeClr":
            return "#" + sch.get(c.get("val"), "FF00FF")
    return None


def fill_and_line(shape, sch):
    sppr = shape._element.find(".//" + P + "spPr")
    fill = line = None
    if sppr is None:
        return None, None
    for c in sppr:
        if c.tag == A + "solidFill":
            fill = colour(c, sch)
        if c.tag == A + "ln" and c.find(A + "noFill") is None:
            line = colour(c, sch)
    return fill, line


def wrap(d, text, x, y, w, pt, font_file, ink):
    font = ImageFont.truetype(font_file, max(8, int(pt)))
    for para in text.split("\n"):
        line = ""
        for word in para.split(" "):
            trial = (line + " " + word).strip()
            if font.getlength(trial) > w and line:
                d.text((x, y), line, fill=ink, font=font)
                y += pt * 1.25
                line = word
            else:
                line = trial
        d.text((x, y), line, fill=ink, font=font)
        y += pt * 1.25


def draw(d, shapes, sch):
    for sh in shapes:
        if sh.left is None:
            continue
        x, y, w, h = [
            Emu(v).inches * DPI for v in (sh.left, sh.top, sh.width, sh.height)
        ]
        if sh.shape_type == MSO_SHAPE_TYPE.TABLE:
            yy = y
            for r, row in enumerate(sh.table.rows):
                xx, rh = x, Emu(row.height).inches * DPI
                for ci, cell in enumerate(row.cells):
                    cw = Emu(sh.table.columns[ci].width).inches * DPI
                    if r == 0:
                        d.rectangle([xx, yy, xx + cw, yy + rh], fill="#" + sch["lt2"])
                    wrap(d, cell.text, xx + 6, yy + 4, cw - 12, 13, SANS, "#222222")
                    xx += cw
                d.line([x, yy + rh, x + w, yy + rh], fill="#CCCCCC")
                yy += rh
            continue
        if sh.has_chart:
            d.rectangle([x, y, x + w, y + h], outline="#BBBBBB")
            d.text((x + 5, y + 5), "[chart]", fill="#999999")
            continue
        fill, line = fill_and_line(sh, sch)
        geom = sh._element.find(".//" + A + "prstGeom")
        kind = geom.get("prst") if geom is not None else "rect"
        if kind == "line":
            d.line([x, y, x + w, y + h], fill=line or "#888888", width=2)
        elif kind == "ellipse":
            d.ellipse([x, y, x + w, y + h], fill=fill, outline=line)
        elif fill or line:
            radius = 6 if "round" in kind.lower() else 0
            box = [x, y, x + max(w, 1), y + max(h, 1)]
            d.rounded_rectangle(box, radius=radius, fill=fill, outline=line)
        if sh.has_text_frame and sh.text_frame.text.strip():
            runs = sh.text_frame.paragraphs[0].runs
            run = runs[0] if runs else None
            pt = run.font.size.pt if run is not None and run.font.size else 14
            ink = "#222222"
            if run is not None:
                ink = colour(run._r.find(".//" + A + "solidFill"), sch) or ink
            bold = run is not None and run.font.bold
            wrap(d, sh.text_frame.text, x, y, w, pt, BOLD if bold else SANS, ink)


def main() -> None:
    deck = Presentation(sys.argv[1])
    out = Path(sys.argv[2] if len(sys.argv) > 2 else "out/preview")
    out.mkdir(parents=True, exist_ok=True)
    sch = scheme(deck)
    size = (
        int(Emu(deck.slide_width).inches * DPI),
        int(Emu(deck.slide_height).inches * DPI),
    )
    pages = []
    for n, slide in enumerate(deck.slides, 1):
        img = Image.new("RGB", size, "white")
        d = ImageDraw.Draw(img)
        bg = colour(slide.slide_layout._element.find(".//" + P + "bg"), sch)
        if bg:
            d.rectangle([0, 0, *size], fill=bg)
        draw(d, [s for s in slide.slide_layout.shapes if not s.is_placeholder], sch)
        draw(d, slide.shapes, sch)
        path = out / f"slide-{n:02d}.png"
        img.save(path)
        pages.append(path)
    w, h = size
    for k in range(0, len(pages), 6):
        sheet = Image.new("RGB", (w * 2 + 10, (h + 10) * 3), "#777777")
        for i, p in enumerate(pages[k : k + 6]):
            sheet.paste(Image.open(p), ((i % 2) * (w + 10), (i // 2) * (h + 10)))
        sheet.save(out / f"sheet-{k // 6}.png")
    print(f"{len(pages)} slides -> {out}")


if __name__ == "__main__":
    main()
