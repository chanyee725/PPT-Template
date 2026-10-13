"""Flag text boxes whose wrapped text is taller than the box.

    python tools/fit.py templates/interim-report.pptx

Measures with the font fc-match finds for "Pretendard:bold" -- the deck's font
when it is installed, a fallback (and so only an estimate) when it is not.
"""

import subprocess
import sys

from PIL import ImageFont
from pptx import Presentation
from pptx.util import Emu

LINE = 1.25  # line height as a multiple of the font size


def font_path(*patterns: str) -> str:
    """The first pattern fc-match really has; it never fails, it substitutes."""
    for pattern in patterns:
        out = subprocess.run(
            ["fc-match", "-f", "%{family}\t%{file}", pattern],
            capture_output=True,
            text=True,
        ).stdout
        family, _, path = out.partition("\t")
        if pattern.split(":")[0].lower() in family.lower():
            return path
    return path


def width(font_file: str, text: str, pt: float) -> float:
    font = ImageFont.truetype(font_file, int(pt * 10))
    return font.getlength(text) / 10 / 72


def needed_height(frame, box_w: float, font_file: str) -> float:
    total = 0.0
    for p in frame.paragraphs:
        pt = max([r.font.size.pt for r in p.runs if r.font.size] or [18])
        bulleted = any(x.tag.endswith("buChar") for x in p._p.iter())
        usable = box_w - (0.3 if bulleted else 0)
        text = "".join(r.text for r in p.runs)
        for segment in text.split("\n") if text else [""]:
            lines, current = 1, ""
            for word in segment.split(" "):
                trial = (current + " " + word).strip()
                if width(font_file, trial, pt) > usable and current:
                    lines, current = lines + 1, word
                else:
                    current = trial
            total += lines * pt * LINE / 72
    return total


def main() -> int:
    deck = Presentation(sys.argv[1])
    font_file = sys.argv[2] if len(sys.argv) > 2 else font_path("Pretendard:bold", "NanumGothic:bold")
    sw, sh = deck.slide_width, deck.slide_height
    problems = 0
    for n, slide in enumerate(deck.slides, 1):
        for shape in slide.shapes:
            if shape.left is None:
                continue
            if (
                shape.left < 0
                or shape.top < 0
                or shape.left + shape.width > sw
                or shape.top + shape.height > sh
            ):
                print(f"slide {n}: {shape.name} runs off the slide")
                problems += 1
            if not shape.has_text_frame or not shape.text_frame.text.strip():
                continue
            if shape.name.startswith("Slide Number"):
                continue
            w, h = Emu(shape.width).inches, Emu(shape.height).inches
            need = needed_height(shape.text_frame, w, font_file)
            if need > h + 0.02:
                print(
                    f"slide {n}: {shape.name} needs {need:.2f}in, box is {h:.2f}in"
                    f" -- {shape.text_frame.text[:40]!r}"
                )
                problems += 1
    print("ok" if not problems else f"{problems} problem(s)")
    return 1 if problems else 0


if __name__ == "__main__":
    sys.exit(main())
